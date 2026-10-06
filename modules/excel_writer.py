"""
excel_writer.py
スキャン結果を Excel ファイルに記録するモジュール
"""

import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


HEADERS = ["No.", "JANコード", "種別", "スキャン日時", "備考"]
HEADER_BG = "4472C4"  # 青


def _get_or_create_workbook(path: str, sheet_name: str):
    """既存ファイルを開くか、新規作成する。"""
    if os.path.exists(path):
        wb = openpyxl.load_workbook(path)
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.create_sheet(sheet_name)
            _write_header(ws)
    else:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        _write_header(ws)
    return wb, ws


def _write_header(ws):
    """ヘッダー行を書き込む。"""
    for col, text in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_BG)
        cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 20


def append_scan(path: str, sheet_name: str, code: str, code_type: str, note: str = ""):
    """
    スキャン結果を Excel に 1 行追記する。

    Args:
        path: Excel ファイルパス
        sheet_name: シート名
        code: バーコード文字列
        code_type: バーコード種別 (EAN13 等)
        note: 備考
    """
    wb, ws = _get_or_create_workbook(path, sheet_name)
    # 現在の最終行を取得（ヘッダー含む）
    next_row = ws.max_row + 1 if ws.max_row > 1 or ws.cell(1, 1).value else 2
    no = next_row - 1  # ヘッダー除いた連番
    ws.cell(next_row, 1, no)
    ws.cell(next_row, 2, code)
    ws.cell(next_row, 3, code_type)
    ws.cell(next_row, 4, datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    ws.cell(next_row, 5, note)
    wb.save(path)
